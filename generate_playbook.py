import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_playbook():
    doc = Document()

    # 1. Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Color Palette Tokens (Devoteam Brand April 2026)
    COLOR_RED_POPPY = RGBColor(248, 72, 94)      # #f8485e
    COLOR_DARK_GREY = RGBColor(60, 60, 58)       # #3c3c3a
    COLOR_WHITE = RGBColor(255, 255, 255)        # #ffffff
    COLOR_AQUA = RGBColor(215, 235, 231)         # #d7ebe7
    COLOR_LIGHT_GREY = RGBColor(239, 238, 238)   # #efeeee
    COLOR_MUTED_GREY = RGBColor(120, 120, 118)

    HEX_DARK_GREY = "3c3c3a"
    HEX_LIGHT_GREY = "efeeee"
    HEX_AQUA = "d7ebe7"
    HEX_RED_POPPY = "f8485e"

    # 3. Typography Styles Setup (Montserrat)
    styles = doc.styles

    # Normal Style
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Montserrat'
    normal_font.size = Pt(10)
    normal_font.color.rgb = COLOR_DARK_GREY

    # Helper function for setting cell shading
    def set_cell_background(cell, hex_color):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    # Helper function for setting cell margins/padding
    def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Helper function to add Devoteam branded callout card
    def add_callout(text, title="OPERATIONAL GUARDRAIL", hex_bg=HEX_LIGHT_GREY, hex_border=HEX_RED_POPPY):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, hex_bg)
        set_cell_margins(cell, top=160, bottom=160, left=200, right=200)

        # Set left border thick and colored
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{hex_border}"/>
                <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
                <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run_title = p.add_run(f"{title}: ")
        run_title.font.name = 'Montserrat'
        run_title.font.size = Pt(9.5)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_RED_POPPY

        run_text = p.add_run(text)
        run_text.font.name = 'Montserrat'
        run_text.font.size = Pt(9.5)
        run_text.font.color.rgb = COLOR_DARK_GREY

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Header: Logo & Eyebrow
    logo_path = "/Users/agungfadlan/.gemini/config/skills/devoteam-brand/assets/logos/devoteam-primary.png"
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_logo.paragraph_format.space_after = Pt(18)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.8))

    # Eyebrow / Category
    p_eye = doc.add_paragraph()
    p_eye.paragraph_format.space_before = Pt(0)
    p_eye.paragraph_format.space_after = Pt(4)
    run_eye = p_eye.add_run("GOOGLE WORKSPACE SEC-OPS PLAYBOOK • STANDARD OPERATING PROCEDURE")
    run_eye.font.name = 'Montserrat'
    run_eye.font.size = Pt(9)
    run_eye.font.bold = True
    run_eye.font.color.rgb = COLOR_RED_POPPY

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("Threat Containment & Email Purge Playbook")
    run_title.font.name = 'Montserrat'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_DARK_GREY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Rapid Mailbox Remediation via GAM / GAMADV-XTD3 and Google Sheets UI Control Center")
    run_sub.font.name = 'Montserrat'
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = COLOR_MUTED_GREY

    # Metadata Control Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document Reference", "DEV-SOP-GWS-PURGE-2026-001"),
        ("Target Service", "Google Workspace (Gmail API & Admin Directory API)"),
        ("Author / Custodian", "Devoteam G Cloud — Architecture & Managed SecOps"),
        ("Classification", "Internal Enterprise Standard Operating Procedure"),
        ("Version & Date", "1.0.0 • October 2026 (Devoteam Brand April 2026 Standards)")
    ]
    for i, (k, v) in enumerate(meta_data):
        cell_k = meta_table.cell(i, 0)
        cell_v = meta_table.cell(i, 1)
        set_cell_background(cell_k, HEX_LIGHT_GREY)
        set_cell_margins(cell_k, top=60, bottom=60, left=120, right=120)
        set_cell_margins(cell_v, top=60, bottom=60, left=120, right=120)

        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_before = Pt(1)
        pk.paragraph_format.space_after = Pt(1)
        rk = pk.add_run(k)
        rk.font.name = 'Montserrat'
        rk.font.size = Pt(9)
        rk.font.bold = True
        rk.font.color.rgb = COLOR_DARK_GREY

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_before = Pt(1)
        pv.paragraph_format.space_after = Pt(1)
        rv = pv.add_run(v)
        rv.font.name = 'Montserrat'
        rv.font.size = Pt(9)
        rv.font.color.rgb = COLOR_DARK_GREY

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Helper for Headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Montserrat'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK_GREY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Montserrat'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_RED_POPPY
        return p

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Montserrat'
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK_GREY
        r_body = p.add_run(text)
        r_body.font.name = 'Montserrat'
        r_body.font.size = Pt(10)
        r_body.font.color.rgb = COLOR_DARK_GREY
        return p

    # --- SECTION 1: EXECUTIVE SUMMARY ---
    add_h1("1. Executive Summary & Purpose")
    add_p("This Standard Operating Procedure (SOP) defines the operational playbook for executing surgical email clawback and containment operations across enterprise Google Workspace tenants. It governs incidents involving active phishing outbreaks, accidental confidential data leakages, and regulatory compliance expungements.")
    add_p("The procedure combines a central Google Sheets Control Center (which provides query sanitisation, blast-radius validation, and audit tracking) with the Google Apps Manager (GAM / GAMADV-XTD3) command-line automation engine.")

    add_callout(
        "Never execute a hard DELETE or soft TRASH operation without first running a DRY_RUN simulation. Evaluate the blast radius to confirm that legitimate communications are not affected.",
        title="CRITICAL SAFETY MANDATE"
    )

    # --- SECTION 2: ARCHITECTURE & TRUST MODEL ---
    add_h1("2. Architectural Framework & Domain-Wide Delegation (DWD)")
    add_p("Standard Google Workspace administrator accounts cannot read or delete messages from individual user mailboxes via ordinary user OAuth 2.0 credentials. To perform domain-wide email remediation, Google Workspace requires Domain-Wide Delegation of Authority (DWD).")

    add_h2("2.1 The Two-Tier Architecture")
    add_p("1. Control Plane (Google Sheets & Apps Script): Acts as the operator interface. It validates query syntax, checks for catastrophic broad filters, holds target recipient lists, and logs an immutable SHA-256 hash of every incident action.")
    add_p("2. Execution Plane (GAM / GAMADV-XTD3): Holds the Service Account private key with Domain-Wide Delegation. For each target mailbox, GAM mints an administrative JWT token setting the 'sub' claim to the user's email, invoking the Gmail API (users.messages.trash or users.messages.delete) directly.")

    add_h2("2.2 Required OAuth 2.0 Scopes")
    add_p("The Service Account Client ID must be authorised in the Google Workspace Admin Console under Security > Access and data control > API controls > Domain-Wide Delegation with the following scopes:")

    # Scopes Table
    table_scopes = doc.add_table(rows=5, cols=2)
    table_scopes.alignment = WD_TABLE_ALIGNMENT.CENTER
    scopes_data = [
        ("OAuth 2.0 Scope", "Operational Justification"),
        ("https://mail.google.com/", "Grants administrative access to search, move to trash, and permanently delete messages."),
        ("https://www.googleapis.com/auth/admin.directory.user.readonly", "Enables GAM to enumerate target users and evaluate domain-wide recipients."),
        ("https://www.googleapis.com/auth/admin.directory.group.readonly", "Allows GAM to expand Google Groups into individual target mailboxes."),
        ("https://www.googleapis.com/auth/spreadsheets.readonly", "Enables direct real-time batch execution from live Google Sheet tabs (gam csv gsheet).")
    ]
    for row_idx, (c0, c1) in enumerate(scopes_data):
        cell0 = table_scopes.cell(row_idx, 0)
        cell1 = table_scopes.cell(row_idx, 1)
        if row_idx == 0:
            set_cell_background(cell0, HEX_DARK_GREY)
            set_cell_background(cell1, HEX_DARK_GREY)
            txt_color = COLOR_WHITE
            is_bold = True
        else:
            set_cell_background(cell0, HEX_LIGHT_GREY if row_idx % 2 == 1 else "ffffff")
            set_cell_background(cell1, HEX_LIGHT_GREY if row_idx % 2 == 1 else "ffffff")
            txt_color = COLOR_DARK_GREY
            is_bold = False

        set_cell_margins(cell0, 80, 80, 120, 120)
        set_cell_margins(cell1, 80, 80, 120, 120)

        p0 = cell0.paragraphs[0]
        r0 = p0.add_run(c0)
        r0.font.name = 'Montserrat'
        r0.font.size = Pt(8.5)
        r0.font.bold = is_bold
        r0.font.color.rgb = txt_color

        p1 = cell1.paragraphs[0]
        r1 = p1.add_run(c1)
        r1.font.name = 'Montserrat'
        r1.font.size = Pt(8.5)
        r1.font.bold = is_bold
        r1.font.color.rgb = txt_color

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 3: STEP-BY-STEP OPERATING PROCEDURE ---
    add_h1("3. Standard Operating Procedure (SOP)")

    add_h2("Phase 1: Incident Triage & Query Formulation")
    add_p("Identify the malicious message attributes using perimeter gateway logs, SecOps telemetry, or user report headers. Formulate an RFC 822 compliant search query containing at least two restrictive anchors:")
    add_p("• Message-ID Anchor: rfc822msgid:<unique-id@attacker.com>", bold_prefix="Primary Filter: ")
    add_p("• Sender & Subject: from:attacker@evil.com subject:\"Overdue Invoice\"", bold_prefix="Alternative Filter: ")
    add_p("• Date Guardrail: after:2026/10/01 (Restricts search window to avoid historical collateral damage).", bold_prefix="Mandatory Bound: ")

    add_h2("Phase 2: Setting up Google Sheet Control Center")
    add_p("1. Navigate to the enterprise Purge Control Center spreadsheet.")
    add_p("2. Open the interactive UI by selecting '⚡ GAM Email Purge' > '📱 Open Purge Dashboard (Sidebar)'.")
    add_p("3. Fill in the Incident Reference ID (e.g. INC-2026-10-092), target scope, and the Gmail RFC 822 query.")
    add_p("4. If the incident affects specific accounts, paste victim email addresses into the 'Target_Mailboxes' tab.")

    add_h2("Phase 3: Simulation & Blast-Radius Audit (DRY_RUN)")
    add_p("1. Set the Purge Action to 'DRY_RUN (Count & List Only)'.")
    add_p("2. Copy the generated dry-run command or execute via gam-purge-runner.sh:")
    add_p("   gam all users print messages query \"rfc822msgid:<...> after:2026/10/01\"")
    add_p("3. Review the output CSV. Confirm that the message count aligns with threat intelligence and that no internal executive communications match the query.")

    add_h2("Phase 4: Dual-Custody Approval & Containment Execution")
    add_p("1. Review the blast-radius impact report with the Incident Commander or Lead SecOps Architect.")
    add_p("2. Determine containment mode:")
    add_p("   • Soft Purge (TRASH) - Recommended: Moves messages to user Trash. Completely isolates threat while preserving a 30-day administrative recovery window.")
    add_p("   • Hard Purge (DELETE) - Exceptional Cases: Permanently expunges messages. Requires setting 'Safety Confirmation' to 'YES' in cell C11 of the sheet.")
    add_p("3. Execute the containment command via CLI runner or live sheet connector:")
    add_p("   gam csv gsheet \"<SHEET_ID>\" \"Target_Mailboxes\" gam user ~Email trash messages query \"<QUERY>\" doit")

    add_h2("Phase 5: Post-Incident Verification & Audit Sign-Off")
    add_p("1. Re-run the simulation query. Confirm matching message count returns exactly 0 in user inboxes.")
    add_p("2. Verify that the command execution details, operator identity, and SHA-256 hash are recorded in the 'Audit_Log' tab.")
    add_p("3. Attach the audit log export to the SecOps incident ticket.")

    # --- SECTION 4: OPERATIONAL GUARDRAILS ---
    add_h1("4. Operational Guardrails & Pattern Blacklist")
    add_p("The Google Apps Script validator strictly blocks unsafe, broad, or malformed queries before command generation. Any query matching the following patterns will be immediately rejected:")

    # Table of Blacklist
    table_bl = doc.add_table(rows=6, cols=2)
    table_bl.alignment = WD_TABLE_ALIGNMENT.CENTER
    bl_data = [
        ("Disallowed Pattern", "Risk Description & Mitigation"),
        ("Empty Query / Whitespace", "Matches 100% of domain emails. Mandatory rejection."),
        ("Wildcard '*'", "Catastrophic blanket match across entire mailbox."),
        ("is:unread / is:read", "Matches normal corporate correspondence. Specific sender/subject anchor required."),
        ("label:inbox / label:sent", "Matches all active inbox or sent messages without filtering."),
        ("Missing Anchors", "Queries without from:, subject:, rfc822msgid:, or date bounds are flagged for review.")
    ]
    for row_idx, (c0, c1) in enumerate(bl_data):
        cell0 = table_bl.cell(row_idx, 0)
        cell1 = table_bl.cell(row_idx, 1)
        if row_idx == 0:
            set_cell_background(cell0, HEX_DARK_GREY)
            set_cell_background(cell1, HEX_DARK_GREY)
            txt_color = COLOR_WHITE
            is_bold = True
        else:
            set_cell_background(cell0, HEX_LIGHT_GREY if row_idx % 2 == 1 else "ffffff")
            set_cell_background(cell1, HEX_LIGHT_GREY if row_idx % 2 == 1 else "ffffff")
            txt_color = COLOR_DARK_GREY
            is_bold = False

        set_cell_margins(cell0, 80, 80, 120, 120)
        set_cell_margins(cell1, 80, 80, 120, 120)

        p0 = cell0.paragraphs[0]
        r0 = p0.add_run(c0)
        r0.font.name = 'Montserrat'
        r0.font.size = Pt(8.5)
        r0.font.bold = is_bold
        r0.font.color.rgb = txt_color

        p1 = cell1.paragraphs[0]
        r1 = p1.add_run(c1)
        r1.font.name = 'Montserrat'
        r1.font.size = Pt(8.5)
        r1.font.bold = is_bold
        r1.font.color.rgb = txt_color

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 5: COMMAND QUICK REFERENCE ---
    add_h1("5. GAM Command Quick Reference Matrix")

    # Command Table
    table_cmd = doc.add_table(rows=4, cols=3)
    table_cmd.alignment = WD_TABLE_ALIGNMENT.CENTER
    cmd_data = [
        ("Target Scope", "Simulation (Dry Run)", "Containment Action (Trash / Delete)"),
        ("Live Google Sheet (Target_Mailboxes)", "gam csv gsheet \"<ID>\" \"Target_Mailboxes\" gam user ~Email print messages query \"<Q>\"", "gam csv gsheet \"<ID>\" \"Target_Mailboxes\" gam user ~Email trash messages query \"<Q>\" doit"),
        ("All Domain Users (Domain-Wide)", "gam all users print messages query \"<Q>\"", "gam all users trash messages query \"<Q>\" doit"),
        ("Single Mailbox", "gam user <USER> print messages query \"<Q>\"", "gam user <USER> trash messages query \"<Q>\" doit")
    ]
    for row_idx, (c0, c1, c2) in enumerate(cmd_data):
        cell0 = table_cmd.cell(row_idx, 0)
        cell1 = table_cmd.cell(row_idx, 1)
        cell2 = table_cmd.cell(row_idx, 2)
        if row_idx == 0:
            set_cell_background(cell0, HEX_DARK_GREY)
            set_cell_background(cell1, HEX_DARK_GREY)
            set_cell_background(cell2, HEX_DARK_GREY)
            txt_color = COLOR_WHITE
            is_bold = True
        else:
            set_cell_background(cell0, HEX_LIGHT_GREY if row_idx % 2 == 1 else "ffffff")
            set_cell_background(cell1, HEX_LIGHT_GREY if row_idx % 2 == 1 else "ffffff")
            set_cell_background(cell2, HEX_LIGHT_GREY if row_idx % 2 == 1 else "ffffff")
            txt_color = COLOR_DARK_GREY
            is_bold = False

        set_cell_margins(cell0, 80, 80, 100, 100)
        set_cell_margins(cell1, 80, 80, 100, 100)
        set_cell_margins(cell2, 80, 80, 100, 100)

        for cell, text in [(cell0, c0), (cell1, c1), (cell2, c2)]:
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Montserrat'
            r.font.size = Pt(8)
            r.font.bold = is_bold
            r.font.color.rgb = txt_color

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # --- SECTION 6: EMERGENCY RESTORATION ---
    add_h1("6. Emergency Rollback & Recovery Runbook")
    add_p("If an operational error or false positive occurs during a Soft Purge (TRASH), messages can be rapidly restored to user inboxes within 30 days without data loss.")
    add_p("To restore messages from Trash back to the Inbox across affected users, run:")
    add_p("gam csv gsheet \"<SHEET_ID>\" \"Target_Mailboxes\" gam user ~Email untrash messages query \"<QUERY>\" doit", bold_prefix="Restoration Command: ")
    add_p("Alternatively, to restore across all users in the tenant:")
    add_p("gam all users untrash messages query \"<QUERY>\" doit", bold_prefix="Domain Restoration: ")

    # Footer note
    section = doc.sections[0]
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("Devoteam G Cloud • AI-driven tech consulting • Confidential")
    f_run.font.name = 'Montserrat'
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = COLOR_MUTED_GREY

    # Save Document
    output_path = "/Users/agungfadlan/.gemini/antigravity/scratch/gws-email-purge/docs/DEV-SOP-GWS-Email-Purge-Playbook.docx"
    doc.save(output_path)
    print(f"Successfully generated playbook: {output_path}")

if __name__ == "__main__":
    create_playbook()
